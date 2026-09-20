"""DOCX text extraction, including tables."""

from __future__ import annotations

import asyncio
from pathlib import Path

from docx import Document as DocxDocument

from app.documents.exceptions import DocumentProcessingError
from app.documents.parsers.base import DocumentParser
from app.documents.schemas import DocumentPage


class DOCXParser(DocumentParser):
    async def parse(self, file_path: Path) -> list[DocumentPage]:
        return await asyncio.to_thread(self._parse, file_path)

    def _parse(self, file_path: Path) -> list[DocumentPage]:
        try:
            document = DocxDocument(file_path)
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return [DocumentPage(page_number=1, text="\n".join(parts))]
        except DocumentProcessingError:
            raise
        except Exception as error:
            raise DocumentProcessingError("Не удалось прочитать DOCX-файл.") from error
