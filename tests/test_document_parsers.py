import asyncio
from pathlib import Path

import fitz
from docx import Document as DocxDocument
import pytest

from app.documents.exceptions import EmptyDocumentError, UnsupportedDocumentError
from app.documents.parsers.docx import DOCXParser
from app.documents.parsers.pdf import PDFParser
from app.documents.processor import detect_file_type, extract_pages


def test_detect_file_type_from_name_and_mime() -> None:
    assert detect_file_type("spec.PDF") == "pdf"
    assert detect_file_type("file.bin", "application/pdf") == "pdf"
    with pytest.raises(UnsupportedDocumentError):
        detect_file_type("scan.djvu")


def test_pdf_parser_extracts_page_text(tmp_path: Path) -> None:
    path = tmp_path / "spec.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Requirement: license")
    document.save(path)
    document.close()

    pages = asyncio.run(PDFParser().parse(path))

    assert pages[0].page_number == 1
    assert "license" in pages[0].text


def test_docx_parser_extracts_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "spec.docx"
    document = DocxDocument()
    document.add_paragraph("Поставщик должен иметь опыт")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Срок"
    table.cell(0, 1).text = "10 дней"
    document.save(path)

    pages = asyncio.run(DOCXParser().parse(path))

    assert "опыт" in pages[0].text
    assert "10 дней" in pages[0].text


def test_extract_pages_rejects_empty_pdf(tmp_path: Path) -> None:
    path = tmp_path / "empty.pdf"
    document = fitz.open()
    document.new_page()
    document.save(path)
    document.close()

    with pytest.raises(EmptyDocumentError):
        asyncio.run(extract_pages(path, "pdf"))
